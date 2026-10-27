"""Export run results as a new, dated tab in a copy of the test-case workbook.

The new tab is a clone of the source tab (styles, merges, column widths, frozen
panes, dropdowns, conditional formatting), so its layout is identical to the
input. Only result-related cells are rewritten:

  Test ID (only when empty) | Actual result | Result | IOS | ANDROID | Note | Tester
  + the summary block above the header (Tested by, Test Date, Build version, counters).

Result values are PASS / FAIL / PENDING. A case is PASS only when its Maestro
flow passed on every platform that was run; anything not verified is PENDING
with the reason in the Note column.
"""
from __future__ import annotations

import copy
import datetime as dt
import os
import re
import shutil
import warnings
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.workbook.views import BookView

from . import template
from .testcases import csv_sheet_title, read_csv_grid
from .util import UserError, fold, info, run, short, warn

PLATFORM_LABEL = {"ios": "iOS", "android": "Android"}
AUTOMATION_TESTER = "Maestro (auto)"
RAN = ("pass", "fail")
RESULT_FIELDS = ("actual", "result", "ios", "android", "evidence", "tester")
NOTE_PREFIX = {
    "manual": "[Manual]",
    "blocked": "[Blocked]",
    "error": "[Automation error]",
    "not_automated": "[Not automated]",
    "missing": "[Not run]",
    "skipped": "[N/A]",
}
DEFAULT_REASON = {
    "missing": "Not included in this run",
    "not_automated": "No Maestro flow for this case yet",
    "skipped": "Not applicable on this platform",
}
SUMMARY_RULES = (
    ("tester", r"^(tested by|tester|executed by|nguoi test)$"),
    ("date", r"^(test date|date|execution date|ngay test)$"),
    ("build", r"^(build|build version|build no|build number|app version|version)$"),
    ("pass", r"^pass(ed)?( cases?)?$"),
    ("fail", r"^fail(ed)?( cases?)?$"),
    ("pending", r"^(unavailab\w*|pending|blocked|skipped|not run|n/?a)( cases?)?$"),
    ("total", r"^(all|all cases|total|total cases)$"),
)
INVALID_TITLE = re.compile(r"[\[\]:*?/\\]")


@dataclass
class Verdict:
    status: str                                   # PASS / FAIL / PENDING
    actual: str = ""
    note: str = ""
    automated: bool = False
    evidence: Dict[str, str] = field(default_factory=dict)  # platform -> path relative to the run dir


# --------------------------------------------------------------------------- verdicts

def _note(platform: str, rec: dict, multi: bool) -> str:
    name = PLATFORM_LABEL.get(platform, platform)
    status = rec.get("status")
    if status == "pass":
        if rec.get("coverage") == "partial":
            return (f"[Auto - partial] {name}: automated checks passed; verify manually: "
                    f"{rec.get('uncovered') or 'see expected result'}")
        text = f"[Auto] Maestro {name}: PASS"
        if rec.get("duration_s"):
            text += f" in {rec['duration_s']:.0f}s"
        if (rec.get("attempts") or 1) > 1:
            text += f" (attempt {rec['attempts']})"
        return text
    if status == "fail":
        text = f"[Auto] Maestro {name}: FAIL"
        if rec.get("failed_step"):
            text += f" at step {rec['failed_step']}"
        return text
    reason = rec.get("message") or DEFAULT_REASON.get(status, "")
    prefix = f"{name}: " if multi else ""
    return f"{prefix}{NOTE_PREFIX.get(status, '[Pending]')} {short(reason, 300)}".strip()


def combine(case: dict, runs: List[dict]) -> Verdict:
    """Merge per-platform records of one case into the sheet verdict."""
    entries: List[Tuple[str, dict]] = []
    for run_doc in runs:
        rec = (run_doc.get("results") or {}).get(case["key"]) or {"status": "missing"}
        entries.append((run_doc.get("platform", "?"), rec))
    multi = len(entries) > 1
    relevant = [(p, r) for p, r in entries if r.get("status") != "skipped"]
    ran = [(p, r) for p, r in relevant if r.get("status") in RAN]
    failed = [(p, r) for p, r in ran if r["status"] == "fail"]
    unverified = [(p, r) for p, r in relevant if r.get("status") not in RAN]
    partial = [(p, r) for p, r in ran if r.get("coverage") == "partial"]

    if failed:
        status = "FAIL"
        actual = "\n".join((f"{PLATFORM_LABEL.get(p, p)}: " if multi else "") + (r.get("actual") or r.get("message") or "Failed")
                           for p, r in failed)
    elif ran and not unverified and not partial:
        status, actual = "PASS", "As expected"
    else:
        status = "PENDING"
        actual = "Automated checks as expected" if ran and not failed and partial else ""

    same_reason = multi and len({(r.get("status"), r.get("message")) for _, r in entries}) == 1
    if same_reason and entries[0][1].get("status") not in RAN:
        note = _note(entries[0][0], entries[0][1], multi=False)      # identical on every platform: say it once
    else:
        note = "\n".join(_note(p, r, multi) for p, r in entries if r.get("status") != "skipped" or not multi)
    return Verdict(
        status=status,
        actual=actual,
        note=note,
        automated=bool(ran),
        evidence={p: r["evidence"] for p, r in entries if r.get("evidence")},
    )


# --------------------------------------------------------------------------- sheet helpers

def unique_title(wb, base: str) -> str:
    base = INVALID_TITLE.sub("-", base).strip("'").strip()[:31] or "Results"
    taken = {fold(name) for name in wb.sheetnames}
    title, n = base, 2
    while fold(title) in taken:
        suffix = f" ({n})"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    return title


def clone_sheet(wb, src):
    """copy_worksheet + the parts openpyxl does not copy (views, panes, dropdowns, colouring)."""
    new = wb.copy_worksheet(src)
    new.sheet_state = "visible"
    for attr in ("zoomScale", "zoomScaleNormal", "showGridLines", "showRowColHeaders", "rightToLeft"):
        setattr(new.sheet_view, attr, getattr(src.sheet_view, attr))
    new.freeze_panes = src.freeze_panes
    for dv in src.data_validations.dataValidation:
        new.add_data_validation(copy.deepcopy(dv))
    for cf in src.conditional_formatting:
        for rule in cf.rules:
            new.conditional_formatting.add(str(cf.sqref), copy.deepcopy(rule))
    if src.auto_filter.ref:
        new.auto_filter.ref = src.auto_filter.ref
    if src.print_title_rows:
        new.print_title_rows = src.print_title_rows
    return new


def activate(wb, ws) -> None:
    # Google Sheets exports have no <bookViews>; without one openpyxl cannot store the
    # active tab and Excel opens the first (often hidden) sheet.
    if not wb.views:
        wb.views.append(BookView())
    for sheet in wb.worksheets:
        for view in sheet.views.sheetView:
            view.tabSelected = False
    ws.sheet_view.tabSelected = True
    wb.active = wb.index(ws)


def unmerge_result_cells(ws, header_row: int, cols: List[int]) -> None:
    """Old evidence links are often merged across several rows; each case now gets its own."""
    for rng in list(ws.merged_cells.ranges):
        if rng.max_row <= header_row or rng.min_col != rng.max_col or rng.min_col not in cols:
            continue
        style = copy.copy(ws.cell(rng.min_row, rng.min_col)._style)
        coords = list(rng.cells)
        ws.unmerge_cells(str(rng))
        for row, col in coords:
            ws.cell(row, col)._style = copy.copy(style)


def allow_pending(ws, result_col: Optional[int]) -> None:
    """Extend an existing PASS/FAIL dropdown on the result column with PENDING."""
    if not result_col:
        return
    for dv in ws.data_validations.dataValidation:
        covers = any(r.min_col <= result_col <= r.max_col for r in dv.sqref.ranges)
        formula = dv.formula1 or ""
        if dv.type == "list" and covers and formula.startswith('"') and "PENDING" not in formula.upper():
            dv.formula1 = formula[:-1] + ',PENDING"'


def _writable(ws, row: int, col: Optional[int]):
    if not col:
        return None
    cell = ws.cell(row, col)
    return None if isinstance(cell, MergedCell) else cell


def _wrap(cell) -> None:
    alignment = copy.copy(cell.alignment)
    alignment.wrap_text = True
    cell.alignment = alignment


def _as_link(cell, target: str) -> None:
    cell.hyperlink = target
    font = copy.copy(cell.font)
    font.color = "FF0000FF"
    font.u = "single"
    cell.font = font


def write_case(ws, case: dict, cols: Dict[str, int], verdict: Verdict, fill_ids: bool, link_root: str,
               links: Optional[Dict[str, str]] = None) -> None:
    row = case["row"]
    for name in RESULT_FIELDS:
        cell = _writable(ws, row, cols.get(name))
        if cell is not None:
            cell.value = None
            cell.hyperlink = None
            if cell.font.u:             # old link styling would leave an underline on an empty cell
                font = copy.copy(cell.font)
                font.u = None
                cell.font = font

    cell = _writable(ws, row, cols.get("test_id"))
    if fill_ids and cell is not None and cell.value in (None, ""):
        cell.value = case["id"]

    cell = _writable(ws, row, cols.get("result"))
    if cell is not None:
        cell.value = verdict.status

    cell = _writable(ws, row, cols.get("actual"))
    if cell is not None and verdict.actual:
        cell.value = verdict.actual
        _wrap(cell)

    cell = _writable(ws, row, cols.get("note"))
    if cell is not None and verdict.note:
        original = str(cell.value).strip() if cell.value not in (None, "") else ""
        cell.value = f"{original}\n{verdict.note}" if original else verdict.note
        _wrap(cell)

    cell = _writable(ws, row, cols.get("tester"))
    if cell is not None and verdict.automated:
        cell.value = AUTOMATION_TESTER

    for platform, rel in sorted(verdict.evidence.items()):
        name = platform if platform in cols else ("evidence" if "evidence" in cols else None)
        cell = _writable(ws, row, cols.get(name)) if name else None
        if cell is None:
            continue
        target = (links or {}).get(rel) or os.path.normpath(os.path.join(link_root, rel)).replace(os.sep, "/")
        if cell.value:
            cell.value = f"{cell.value}\n{target}"
        else:
            cell.value = target
            _as_link(cell, target)


def _value_cell_right_of(ws, row: int, col: int):
    end = col
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            end = rng.max_col
            break
    cell = ws.cell(row, end + 1)
    if not isinstance(cell, MergedCell):
        return cell
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= end + 1 <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col)
    return None


def update_summary(ws, header_row: int, cols: Dict[str, int], first: int, last: int,
                   tester: str, build: str, run_date: dt.date) -> List[str]:
    """Fill the block above the header (Tested by / Test Date / Build / counters)."""
    res_col = cols.get("result")
    desc_col = cols.get("description") or cols.get("steps")
    res = f"{get_column_letter(res_col)}{first}:{get_column_letter(res_col)}{last}" if res_col else None
    desc = f"{get_column_letter(desc_col)}{first}:{get_column_letter(desc_col)}{last}" if desc_col else None
    values = {
        "tester": tester,
        "date": dt.datetime(run_date.year, run_date.month, run_date.day),
        "build": build or None,
        "pass": f'=COUNTIF({res},"PASS")' if res else None,
        "fail": f'=COUNTIF({res},"FAIL")' if res else None,
        "pending": f'=COUNTIF({res},"PENDING")' if res else None,
        "total": f"=COUNTA({desc})" if desc else None,
    }
    used, updated = set(), []
    for row in range(1, header_row):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            if (row, col) in used or isinstance(cell, MergedCell) or not isinstance(cell.value, str):
                continue
            label = fold(cell.value).rstrip(":").strip()
            kind = next((k for k, pattern in SUMMARY_RULES if re.search(pattern, label)), None)
            if not kind or values.get(kind) is None:
                continue
            target = _value_cell_right_of(ws, row, col)
            if target is None:
                continue
            target.value = values[kind]
            if kind == "date" and target.number_format in ("General", "@"):
                target.number_format = "mmm dd, yyyy"
            used.add((target.row, target.column))
            updated.append(f"{cell.value.strip()} -> {target.coordinate}")
    return updated


# --------------------------------------------------------------------------- export

def default_build(runs: List[dict]) -> str:
    parts = []
    for run_doc in runs:
        app = run_doc.get("app") or {}
        version = app.get("version") or "?"
        build = f" ({app['build']})" if app.get("build") else ""
        parts.append(f"{PLATFORM_LABEL.get(run_doc.get('platform'), run_doc.get('platform'))} {version}{build}")
    return " / ".join(parts)


def run_date_of(runs: List[dict]) -> dt.date:
    for run_doc in runs:
        stamp = run_doc.get("started_at")
        if stamp:
            try:
                return dt.datetime.fromisoformat(stamp).date()
            except ValueError:
                continue
    return dt.date.today()


def _open_source(cases_doc: dict, input_path: Path):
    src = cases_doc["source"]
    suffix = input_path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = load_workbook(input_path)
        if src["sheet"] not in wb.sheetnames:
            raise UserError(f"Sheet '{src['sheet']}' no longer exists in {input_path.name}. Re-run `rnsmoke cases`.")
        return wb, wb[src["sheet"]]
    if suffix in (".csv", ".tsv", ".txt"):
        wb = template.build_from_grid(read_csv_grid(input_path), csv_sheet_title(input_path))
        return wb, wb.active
    raise UserError(f"Cannot export into '{suffix}' files; use the original .xlsx or .csv.")


def _check_layout(ws, cases_doc: dict) -> None:
    src = cases_doc["source"]
    for name, col in src["columns"].items():
        expected = fold(src["headers"].get(name))
        actual = fold(ws.cell(src["header_row"], col).value)
        if expected != actual:
            raise UserError(f"The test-case sheet changed since it was parsed (column {get_column_letter(col)} "
                            f"header is '{ws.cell(src['header_row'], col).value}'). Re-run `rnsmoke cases`.")


def convert_to_xls(path: Path) -> Optional[Path]:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    mac_app = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if not soffice and mac_app.exists():
        soffice = str(mac_app)
    if not soffice:
        warn("LibreOffice not found - skipped .xls conversion (brew install --cask libreoffice).")
        return None
    res = run([soffice, "--headless", "--convert-to", "xls", "--outdir", str(path.parent), str(path)], timeout=240)
    out = path.with_suffix(".xls")
    if res.returncode != 0 or not out.exists():
        warn(f"LibreOffice conversion failed: {short(res.stderr or res.stdout, 300)}")
        return None
    return out


def export(cases_doc: dict, runs: List[dict], run_dir: Path, input_path: Optional[Path] = None,
           out_path: Optional[Path] = None, tab_name: Optional[str] = None, tester: Optional[str] = None,
           build: Optional[str] = None, run_date: Optional[dt.date] = None, fill_ids: bool = True,
           only_new_tab: bool = False, also_xls: bool = False, links: Optional[Dict[str, str]] = None) -> dict:
    """`links` maps run-relative evidence paths to public URLs (see upload.py); others stay local."""
    if not runs:
        raise UserError(f"No results-*.json found in {run_dir}. Run `rnsmoke run` first.")
    src = cases_doc["source"]
    input_path = Path(input_path or src["path"]).expanduser().resolve()
    if not input_path.exists():
        raise UserError(f"Original test-case file not found: {input_path} (pass --input)")
    run_date = run_date or run_date_of(runs)
    tab = tab_name or run_date.strftime("%d%m")
    if out_path is None:
        stem = input_path.stem
        out_path = run_dir / f"{stem} - {tab}.xlsx"
    out_path = Path(out_path).expanduser().resolve()

    wb, src_ws = _open_source(cases_doc, input_path)
    _check_layout(src_ws, cases_doc)
    new_ws = clone_sheet(wb, src_ws)
    new_ws.title = unique_title(wb, tab)
    activate(wb, new_ws)

    cols: Dict[str, int] = src["columns"]
    header_row = src["header_row"]
    unmerge_result_cells(new_ws, header_row, [cols[f] for f in RESULT_FIELDS + ("note",) if f in cols])
    allow_pending(new_ws, cols.get("result"))

    link_root = os.path.relpath(run_dir.resolve(), out_path.parent)
    counts: Counter = Counter()
    evidence_links = Counter()
    for case in cases_doc["cases"]:
        verdict = combine(case, runs)
        write_case(new_ws, case, cols, verdict, fill_ids, link_root, links)
        counts[verdict.status] += 1
        for rel in verdict.evidence.values():
            evidence_links["uploaded" if (links or {}).get(rel) else "local"] += 1

    rows = [c["row"] for c in cases_doc["cases"]]
    summary = update_summary(new_ws, header_row, cols, min(rows), max(rows),
                             tester or "Maestro automation", build if build is not None else default_build(runs),
                             run_date)
    if only_new_tab:
        for ws in list(wb.worksheets):
            if ws is not new_ws:
                wb.remove(ws)
        activate(wb, new_ws)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    xls = convert_to_xls(out_path) if also_xls else None
    return {
        "output": str(out_path),
        "xls": str(xls) if xls else None,
        "sheet": new_ws.title,
        "source_sheet": src["sheet"],
        "counts": dict(counts),
        "evidence_links": dict(evidence_links),
        "summary_cells": summary,
    }


def verify(path: Path, sheet: str, cases_doc: dict) -> Dict[str, int]:
    """Re-open the written workbook and count the result values in the new tab."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(path, read_only=True)
    ws = wb[sheet]
    col = cases_doc["source"]["columns"].get("result")
    counts: Counter = Counter()
    if col:
        wanted = {c["row"] for c in cases_doc["cases"]}
        for idx, row in enumerate(ws.iter_rows(min_col=col, max_col=col, values_only=True), start=1):
            if idx in wanted:
                counts[str(row[0])] += 1
    wb.close()
    return dict(counts)
