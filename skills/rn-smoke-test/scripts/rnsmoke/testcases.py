"""Parse test-case sheets (.xlsx / .csv) into a normalised "cases" document.

The parser finds the header row (Section / Description / Expected result ...),
maps columns by header name, and returns one case per non-empty row together
with the sheet coordinates needed to write results back into the same layout.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import warnings
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from .util import UserError, as_text, clean, fold, now_iso

# Folded (lower-case, accent-free) header spellings -> canonical field name.
HEADER_ALIASES: Dict[str, Tuple[str, ...]] = {
    "test_id": ("test id", "testid", "test case id", "tc id", "tcid", "case id", "id", "no", "no.", "#", "stt",
                "ma tc", "ma test case"),
    "platform": ("platform", "platforms", "os", "nen tang"),
    "section": ("section", "module", "module name", "feature", "function", "screen", "area", "category",
                "component", "chuc nang", "man hinh"),
    "description": ("description", "test case", "test cases", "test case description", "test case name",
                    "scenario", "test scenario", "title", "summary", "objective", "mo ta", "noi dung",
                    "ten test case"),
    "precondition": ("precondition", "pre condition", "pre-condition", "preconditions", "pre-conditions",
                     "dieu kien", "dieu kien tien quyet"),
    "steps": ("steps", "step", "test steps", "steps to reproduce", "procedure", "cac buoc", "buoc thuc hien"),
    "expected": ("expected result", "expected results", "expected", "expected output", "expected behavior",
                 "expected behaviour", "ket qua mong doi", "ket qua mong muon"),
    "actual": ("actual result", "actual results", "actual", "actual output", "ket qua thuc te"),
    "result": ("result", "status", "pass/fail", "test result", "test status", "ket qua", "trang thai"),
    "ios": ("ios", "evidence ios", "ios evidence", "evidence for ios", "ios result"),
    "android": ("android", "evidence android", "android evidence", "evidence for android", "android result"),
    "evidence": ("evidence", "evidences", "attachment", "attachments", "screenshot", "screenshots", "video",
                 "link", "bang chung"),
    "test_data": ("test data", "data", "du lieu", "du lieu test"),
    "note": ("note", "notes", "comment", "comments", "remark", "remarks", "ghi chu"),
    "tester": ("tester", "tested by", "executed by", "assignee", "owner", "nguoi test"),
}

# Fallback keyword rules, applied in order when no alias matches exactly.
KEYWORD_RULES: Tuple[Tuple[str, str], ...] = (
    ("expected", r"\bexpected\b|mong doi|mong muon"),
    ("actual", r"\bactual\b|thuc te"),
    ("ios", r"\bios\b"),
    ("android", r"\bandroid\b"),
    ("test_id", r"\b(test ?id|tc ?id|case ?id)\b"),
    ("description", r"\bdescription\b|\bscenario\b"),
    ("precondition", r"\bpre ?-?conditions?\b"),
    ("steps", r"\bsteps?\b"),
    ("test_data", r"\btest data\b"),
    ("note", r"\bnotes?\b|\bremarks?\b|\bcomments?\b"),
    ("result", r"\bresult\b|\bstatus\b"),
)

TEXT_FIELDS = ("platform", "section", "description", "precondition", "steps", "expected", "test_data", "note")
MAX_HEADER_SCAN = 40
EMPTY_ROWS_TO_STOP = 50


def header_field(value: Any) -> Optional[str]:
    """Return the canonical field for a header cell, or None."""
    text = fold(value)
    if not text:
        return None
    text = re.sub(r"\(.*?\)", " ", text)
    text = re.sub(r"[^a-z0-9/# ]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return None
    for name, aliases in HEADER_ALIASES.items():
        if text in aliases:
            return name
    for name, pattern in KEYWORD_RULES:
        if re.search(pattern, text):
            return name
    return None


def map_header_row(values: List[Any]) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for col, value in enumerate(values, start=1):
        name = header_field(value)
        if name and name not in mapping:
            mapping[name] = col
    return mapping


def parse_platforms(value: Any) -> List[str]:
    text = fold(value)
    platforms = []
    if re.search(r"\bios\b|iphone|ipad", text):
        platforms.append("ios")
    if "android" in text:
        platforms.append("android")
    return platforms or ["ios", "android"]


def case_key(section: str, description: str, expected: str) -> str:
    """Content hash used to link a case to its Maestro flow across sheet revisions."""
    base = "\x1f".join(fold(x) for x in (section, description, expected))
    return hashlib.sha1(base.encode("utf-8")).hexdigest()[:10]


# --------------------------------------------------------------------------- grid

@dataclass
class Grid:
    """A rectangular view over a sheet or CSV with 1-based coordinates."""

    rows: List[List[Any]]
    merged: Dict[Tuple[int, int], Any] = field(default_factory=dict)
    result_col_hint: Optional[int] = None

    @property
    def max_row(self) -> int:
        return len(self.rows)

    @property
    def max_col(self) -> int:
        return max((len(r) for r in self.rows), default=0)

    def value(self, row: int, col: Optional[int]) -> Any:
        if not col or row < 1 or row > len(self.rows):
            return None
        cells = self.rows[row - 1]
        return clean(cells[col - 1]) if col <= len(cells) else None

    def merged_value(self, row: int, col: Optional[int]) -> Any:
        if not col:
            return None
        return clean(self.merged.get((row, col)))

    def row_values(self, row: int) -> List[Any]:
        return [self.value(row, c) for c in range(1, self.max_col + 1)]


def grid_from_sheet(ws, max_rows: Optional[int] = None) -> Grid:
    rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=max_rows or ws.max_row, values_only=True)]
    merged: Dict[Tuple[int, int], Any] = {}
    for rng in ws.merged_cells.ranges:
        anchor = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    merged[(r, c)] = anchor
    hint = None
    for dv in ws.data_validations.dataValidation:
        if dv.type == "list" and dv.formula1 and "PASS" in str(dv.formula1).upper():
            hint = min(r.min_col for r in dv.sqref.ranges)
            break
    return Grid(rows=rows, merged=merged, result_col_hint=hint)


def read_csv_grid(path: Path) -> Grid:
    raw = path.read_bytes()
    text = None
    candidates = ["utf-8-sig"]
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        candidates.insert(0, "utf-16")
    candidates += ["cp1258", "cp1252", "latin-1"]
    for encoding in candidates:
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    first_line = text.split("\n", 1)[0] if text else ""
    delimiter = max((",", ";", "\t"), key=first_line.count)
    rows = [list(r) for r in csv.reader(io.StringIO(text or ""), delimiter=delimiter)]
    return Grid(rows=rows)


def csv_sheet_title(path: Path) -> str:
    """Google Sheets names CSV downloads "<Spreadsheet> - <Tab>.csv"; reuse the tab name."""
    stem = path.stem
    if " - " in stem:
        tail = stem.rsplit(" - ", 1)[1].strip()
        if 0 < len(tail) <= 31:
            return tail
    return "Test cases"


# --------------------------------------------------------------------------- parsing

def detect_header(grid: Grid) -> Tuple[int, Dict[str, int]]:
    best: Tuple[int, Dict[str, int]] = (0, {})
    for row in range(1, min(grid.max_row, MAX_HEADER_SCAN) + 1):
        mapping = map_header_row(grid.row_values(row))
        has_core = ("description" in mapping or "steps" in mapping) and "expected" in mapping
        if has_core and len(mapping) >= 3:
            return row, mapping
        if len(mapping) > len(best[1]):
            best = (row, mapping)
    if len(best[1]) >= 4 and ("description" in best[1] or "steps" in best[1]):
        return best
    raise UserError("Could not find a header row with at least 'Description' and 'Expected result' columns "
                    f"in the first {MAX_HEADER_SCAN} rows.")


def parse_grid(grid: Grid) -> Tuple[int, Dict[str, int], List[dict]]:
    header_row, mapping = detect_header(grid)

    if "result" not in mapping:
        hint = grid.result_col_hint
        after_actual = mapping["actual"] + 1 if "actual" in mapping else None
        if hint and hint not in mapping.values():
            mapping["result"] = hint
        elif after_actual and after_actual not in mapping.values() and not grid.value(header_row, after_actual):
            mapping["result"] = after_actual

    desc_field = "description" if "description" in mapping else "steps"
    cases: List[dict] = []
    seen_keys: Dict[str, int] = {}
    seen_ids: set = set()
    last_section: Optional[str] = None
    empty_run = 0

    for row in range(header_row + 1, grid.max_row + 1):
        rec = {name: grid.value(row, col) for name, col in mapping.items()}
        for grouped in ("section", "platform"):
            if rec.get(grouped) is None:
                rec[grouped] = grid.merged_value(row, mapping.get(grouped))
        description = as_text(rec.get(desc_field))
        expected = as_text(rec.get("expected"))
        if not description and not expected:
            if any(v is not None for v in rec.values()):
                empty_run = 0
                if rec.get("section"):
                    last_section = as_text(rec["section"])
            else:
                empty_run += 1
                if empty_run >= EMPTY_ROWS_TO_STOP:
                    break
            continue
        empty_run = 0

        section = as_text(rec.get("section")) or last_section or "General"
        last_section = section
        number = len(cases) + 1
        sheet_id = as_text(rec.get("test_id"))
        case_id = sheet_id or f"TC-{number:03d}"
        if case_id in seen_ids:
            case_id = f"{case_id}-{number}"
        seen_ids.add(case_id)

        key = case_key(section, description, expected)
        if key in seen_keys:
            seen_keys[key] += 1
            key = f"{key}-{seen_keys[key]}"
        else:
            seen_keys[key] = 1

        case = {
            "id": case_id,
            "sheet_test_id": sheet_id or None,
            "key": key,
            "row": row,
            "section": section,
            "description": description,
            "expected": expected,
            "platforms": parse_platforms(rec.get("platform")),
        }
        for name in ("platform", "precondition", "steps", "test_data", "note"):
            if name in mapping and name != desc_field:
                case[name] = as_text(rec.get(name))
        if "result" in mapping:
            case["previous_result"] = as_text(rec.get("result"))
        cases.append(case)

    if not cases:
        raise UserError(f"Header found on row {header_row} but no test-case rows below it.")
    return header_row, mapping, cases


def choose_sheet(wb, requested: Optional[str] = None):
    if requested:
        if requested not in wb.sheetnames:
            raise UserError(f"Sheet '{requested}' not found. Available: {', '.join(wb.sheetnames)}")
        return wb[requested], []
    candidates = []
    for ws in wb.worksheets:
        try:
            detect_header(grid_from_sheet(ws, max_rows=MAX_HEADER_SCAN))
            candidates.append(ws)
        except UserError:
            continue
    if not candidates:
        raise UserError("No sheet contains a test-case header (Section / Description / Expected result).")
    visible = [ws for ws in candidates if ws.sheet_state == "visible"]
    if wb.active in visible:
        chosen = wb.active
    elif visible:
        chosen = visible[-1]
    else:
        chosen = candidates[-1]
    return chosen, [ws.title for ws in candidates]


def load_cases(path: Path, sheet: Optional[str] = None) -> dict:
    """Parse a test-case file and return the cases document (see module docstring)."""
    path = Path(path).expanduser().resolve()
    if not path.exists():
        raise UserError(f"Test-case file not found: {path}")
    suffix = path.suffix.lower()
    candidates: List[str] = []
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = load_workbook(path, data_only=True)
        ws, candidates = choose_sheet(wb, sheet)
        grid = grid_from_sheet(ws)
        sheet_name, kind = ws.title, "xlsx"
    elif suffix in (".csv", ".tsv", ".txt"):
        grid = read_csv_grid(path)
        sheet_name, kind = csv_sheet_title(path), "csv"
    elif suffix == ".xls":
        raise UserError("Legacy .xls input is not supported. Open it in Excel / Google Sheets and save as .xlsx "
                        "(or export the tab as .csv).")
    else:
        raise UserError(f"Unsupported test-case file type '{suffix}'. Use .xlsx or .csv.")

    header_row, mapping, cases = parse_grid(grid)
    return {
        "version": 1,
        "parsed_at": now_iso(),
        "source": {
            "path": str(path),
            "type": kind,
            "sheet": sheet_name,
            "candidate_sheets": candidates,
            "header_row": header_row,
            "columns": mapping,
            "headers": {name: as_text(grid.value(header_row, col)) for name, col in mapping.items()},
            "first_row": cases[0]["row"],
            "last_row": cases[-1]["row"],
        },
        "cases": cases,
    }


def summarize(doc: dict) -> str:
    src = doc["source"]
    lines = [
        f"Source : {src['path']}",
        f"Sheet  : {src['sheet']}  (header row {src['header_row']}, rows {src['first_row']}-{src['last_row']})",
    ]
    if len(src.get("candidate_sheets") or []) > 1:
        lines.append(f"         other sheets with a test-case header: "
                     f"{', '.join(s for s in src['candidate_sheets'] if s != src['sheet'])}")
    lines.append("Columns: " + ", ".join(f"{k}={src['headers'].get(k) or '?'}" for k in src["columns"]))
    missing = [f for f in ("section", "description", "expected", "result") if f not in src["columns"]]
    if missing:
        lines.append(f"Missing columns: {', '.join(missing)}")
    counts: Dict[str, int] = {}
    for case in doc["cases"]:
        counts[case["section"]] = counts.get(case["section"], 0) + 1
    lines.append(f"Cases  : {len(doc['cases'])} in {len(counts)} sections")
    for section, n in counts.items():
        lines.append(f"  {n:4d}  {section}")
    return "\n".join(lines)
