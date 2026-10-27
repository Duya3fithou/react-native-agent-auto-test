"""Default smoke-checklist look and workbook builders.

Used when the input carries no formatting of its own (a CSV export, or test
cases generated from source code). The look mirrors the team's Google Sheets
checklist: blue summary labels, a grey header row, a PASS/FAIL/PENDING dropdown
and green/red result colouring.
"""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .testcases import Grid, detect_header
from .util import UserError, clean, fold

STATUS_VALUES = ("PASS", "FAIL", "PENDING")
BLUE, GREY, WHITE = "FF3C78D8", "FFB7B7B7", "FFFFFFFF"
GREEN, RED, YELLOW = "FFB7E1CD", "FFE06666", "FFFFE599"
FONT_NAME, FONT_SIZE = "Arial", 10
_THIN = Side(style="thin", color="FF000000")
BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

STANDARD_COLUMNS = (
    ("test_id", "Test ID"), ("platform", "Platform"), ("section", "Section"), ("description", "Description"),
    ("expected", "Expected result"), ("actual", "Actual result"), ("result", "Result"), ("ios", "IOS"),
    ("android", "ANDROID"), ("test_data", "Test data"), ("note", "Note"), ("tester", "Tester"),
)
COLUMN_WIDTHS = {
    "test_id": 10, "platform": 19, "section": 21, "description": 42, "precondition": 30, "steps": 42,
    "expected": 56, "actual": 32, "result": 11, "ios": 26, "android": 26, "evidence": 26, "test_data": 22,
    "note": 32, "tester": 16,
}
WRAP_FIELDS = {"section", "description", "precondition", "steps", "expected", "actual", "test_data", "note"}
SUMMARY_LABELS = {
    "name", "test cases version", "version", "sprint", "build version", "build", "pre-condition", "precondition",
    "pass cases", "fail cases", "unavailable cases", "unavailabe cases", "pending cases", "all cases", "total",
    "tested by", "test date",
}
STANDARD_HEADER_ROW = 10
DEFAULT_PLATFORM = "App / iOS / ANDROID"


def label_style(cell) -> None:
    cell.font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", start_color=BLUE, end_color=BLUE)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    cell.border = BOX


def value_style(cell) -> None:
    cell.font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    cell.border = BOX


def header_style(cell) -> None:
    cell.font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True)
    cell.fill = PatternFill("solid", start_color=GREY, end_color=GREY)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def data_style(cell, wrap: bool) -> None:
    cell.font = Font(name=FONT_NAME, size=FONT_SIZE)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=wrap)


def apply_result_rules(ws, col: int, first_row: int, last_row: int) -> None:
    """Status dropdown + PASS/FAIL/PENDING colouring on the result column."""
    letter = get_column_letter(col)
    ref = f"{letter}{first_row}:{letter}{last_row}"
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_VALUES) + '"', allow_blank=True)
    dv.add(ref)
    ws.add_data_validation(dv)
    for value, color in (("PASS", GREEN), ("FAIL", RED), ("PENDING", YELLOW)):
        ws.conditional_formatting.add(
            ref, CellIsRule(operator="equal", formula=[f'"{value}"'],
                            fill=PatternFill("solid", start_color=color, end_color=color)))


def _coerce(value: Any) -> Any:
    value = clean(value)
    if isinstance(value, str) and value.isdigit() and len(value) < 10 and not value.startswith("0"):
        return int(value)
    return value


def _apply_columns(ws, mapping: Dict[str, int], max_col: int) -> None:
    by_col = {col: name for name, col in mapping.items()}
    for col in range(1, max_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = COLUMN_WIDTHS.get(by_col.get(col, ""), 14)


def build_from_grid(grid: Grid, title: str) -> Workbook:
    """Recreate a CSV export as a styled sheet, cell-for-cell at the same coordinates."""
    header_row, mapping = detect_header(grid)
    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Test cases")[:31]
    max_col = max(grid.max_col, len(STANDARD_COLUMNS))

    for r in range(1, grid.max_row + 1):
        for c in range(1, grid.max_col + 1):
            value = _coerce(grid.value(r, c))
            if value is not None:
                ws.cell(r, c, value)

    for r in range(1, header_row):
        filled = [c for c in range(1, grid.max_col + 1) if grid.value(r, c) is not None]
        if not filled:
            continue
        a_is_label = grid.value(r, 1) is not None
        for c in filled:
            text = fold(grid.value(r, c)).rstrip(":").strip()
            if c == 1 or text in SUMMARY_LABELS:
                label_style(ws.cell(r, c))
            else:
                value_style(ws.cell(r, c))
        if a_is_label and grid.value(r, 3) is None and grid.value(r, 4) is None:
            for c in (2, 3, 4):
                value_style(ws.cell(r, c))
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        elif a_is_label and grid.value(r, 2) is None:
            value_style(ws.cell(r, 2))

    for c in range(1, grid.max_col + 1):
        if grid.value(header_row, c) is not None:
            header_style(ws.cell(header_row, c))

    wrap_cols = {mapping[f] for f in WRAP_FIELDS if f in mapping}
    for r in range(header_row + 1, grid.max_row + 1):
        for c in range(1, grid.max_col + 1):
            data_style(ws.cell(r, c), c in wrap_cols)

    _apply_columns(ws, mapping, max_col)
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    if "result" in mapping:
        apply_result_rules(ws, mapping["result"], header_row + 1, max(grid.max_row, header_row + 500))
    return wb


def _standard_rows(title: str, version: str) -> List[List[Any]]:
    return [
        ["Name", title],
        ["Test cases Version", version],
        ["Sprint", None, "All Cases", None],
        ["Build version"],
        ["Pre-condition"],
        ["Pass Cases", None, "Unavailable Cases", None],
        [None, None, "Fail Cases", None],
        ["Tested by"],
        ["Test Date"],
        [label for _, label in STANDARD_COLUMNS],
    ]


def _case_row(case: dict) -> List[Any]:
    values = {
        "test_id": case.get("id") or case.get("test_id"),
        "platform": case.get("platform") or DEFAULT_PLATFORM,
        "section": case.get("section"),
        "description": case.get("description"),
        "expected": case.get("expected") or case.get("expected_result"),
        "test_data": case.get("test_data"),
        "note": case.get("note"),
    }
    return [values.get(name) for name, _ in STANDARD_COLUMNS]


def build_standard(title: str, cases: List[dict], version: str = "v1.0") -> Workbook:
    """A brand-new checklist in the team's standard layout (summary block + header on row 10)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Test cases"
    rows = _standard_rows(title, version)
    for r, values in enumerate(rows, start=1):
        for c, value in enumerate(values, start=1):
            if value is not None:
                ws.cell(r, c, value)
    first = STANDARD_HEADER_ROW + 1
    for i, case in enumerate(cases):
        for c, value in enumerate(_case_row(case), start=1):
            if clean(value) is not None:
                ws.cell(first + i, c, clean(value))
    last = max(first + len(cases) - 1, first)
    cols = {name: idx for idx, (name, _) in enumerate(STANDARD_COLUMNS, start=1)}
    res, desc = get_column_letter(cols["result"]), get_column_letter(cols["description"])

    for r in range(1, STANDARD_HEADER_ROW):
        values = rows[r - 1]
        if len(values) >= 3 and values[2] is not None:
            label_style(ws.cell(r, 3))
            value_style(ws.cell(r, 4))
        if values[0] is None:
            continue
        label_style(ws.cell(r, 1))
        if len(values) <= 2:
            for c in (2, 3, 4):
                value_style(ws.cell(r, c))
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        else:
            value_style(ws.cell(r, 2))
    ws["D3"] = f"=COUNTA({desc}{first}:{desc}{last})"
    ws["B6"] = f'=COUNTIF({res}{first}:{res}{last},"PASS")'
    ws["D6"] = f'=COUNTIF({res}{first}:{res}{last},"PENDING")'
    ws["D7"] = f'=COUNTIF({res}{first}:{res}{last},"FAIL")'
    ws["B9"].number_format = "mmm dd, yyyy"

    for c in range(1, len(STANDARD_COLUMNS) + 1):
        header_style(ws.cell(STANDARD_HEADER_ROW, c))
    wrap_cols = {cols[f] for f in WRAP_FIELDS if f in cols}
    for r in range(first, last + 1):
        for c in range(1, len(STANDARD_COLUMNS) + 1):
            data_style(ws.cell(r, c), c in wrap_cols)
    _apply_columns(ws, cols, len(STANDARD_COLUMNS))
    ws.freeze_panes = ws.cell(first, 1)
    apply_result_rules(ws, cols["result"], first, max(last, first + 500))
    return wb


def write_template(path: Path, title: str, cases: Optional[List[dict]] = None, version: str = "v1.0") -> Path:
    """Write a checklist (blank or pre-filled) as .xlsx or .csv in the standard layout."""
    path = Path(path).expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    cases = cases or []
    for i, case in enumerate(cases, start=1):
        if not case.get("description") or not (case.get("expected") or case.get("expected_result")):
            raise UserError(f"Generated case #{i} needs both 'description' and 'expected'.")
        case.setdefault("section", "General")
    if path.suffix.lower() == ".csv":
        with path.open("w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.writer(fh)
            for row in _standard_rows(title, version):
                writer.writerow(list(row) + [None] * (len(STANDARD_COLUMNS) - len(row)))
            for case in cases:
                writer.writerow(_case_row(case))
    elif path.suffix.lower() == ".xlsx":
        build_standard(title, cases, version).save(path)
    else:
        raise UserError("Template output must end with .xlsx or .csv")
    return path
