"""Unit tests for the rnsmoke package (no device or Maestro needed).

Run:  python3 -m unittest discover -s tests -v
"""
from __future__ import annotations

import csv
import datetime as dt
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills" / "rn-smoke-test" / "scripts"))

import io  # noqa: E402
import os  # noqa: E402
import urllib.error  # noqa: E402
from unittest import mock  # noqa: E402

from rnsmoke import gdrive, report, runner, safety, template, testcases, upload  # noqa: E402
from rnsmoke.util import UserError  # noqa: E402
from rnsmoke.workspace import Workspace, select_cases  # noqa: E402

HEADER = ["Test ID", "Platform", "Section", "Description", "Expected result", "Actual result", "Result",
          "IOS", "ANDROID", "Test data", "Note ", "Tester"]
CASES = [
    ("App / iOS / ANDROID", "Registration", "Register without first and last name",
     '"First name is required" popup appears\n"Last name is required" popup appears'),
    ("App / iOS / ANDROID", None, "Register with mismatching passwords", '"Passwords do not match" is displayed'),
    ("App / iOS / ANDROID", "Login", "Login with valid credentials", "Home screen is displayed"),
    ("iOS", "Login", "Login with a wrong password", '"Invalid credentials" popup appears'),
    ("ANDROID", "Home", "Home screen for a user with 2 profiles", "1. Facility switcher\n2. Profile switcher"),
    ("App / iOS / ANDROID", "Home", "Switch facility", "A list of linked facilities is displayed"),
]
SUMMARY = [
    ["Name", "Demo Smoke Test Checklist"],
    ["Test cases Version", "v1.0 Checklist"],
    ["Sprint", None, "All Cases", "=COUNT(A12:A1782)"],
    ["Build version"],
    ["Pre-condition "],
    ["Pass Cases ", '=COUNTIF(G2:G1782, "Pass")', "Unavailabe Cases", '=COUNTIF(H2:H1782, "Unavailable")'],
    [None, None, "Fail Cases", '=COUNTIF(H2:H1782, "Fail")'],
    ["Tested by", "QA"],
    ["Test Date", dt.datetime(2026, 4, 2)],
]


def make_sheet(ws, with_result_header=True, dv_values='"PASS,FAIL"'):
    for r, values in enumerate(SUMMARY, start=1):
        for c, value in enumerate(values, start=1):
            if value is not None:
                ws.cell(r, c, value)
        if values and values[0] is not None:
            ws.cell(r, 1).fill = PatternFill("solid", start_color="FF3C78D8", end_color="FF3C78D8")
            ws.cell(r, 1).font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
    for r in (1, 2, 4, 5, 8, 9):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    for c, value in enumerate(HEADER, start=1):
        if value == "Result" and not with_result_header:
            continue
        ws.cell(10, c, value)
    for i, (platform, section, desc, expected) in enumerate(CASES):
        row = 11 + i
        ws.cell(row, 2, platform)
        if section:
            ws.cell(row, 3, section)
        ws.cell(row, 4, desc)
        ws.cell(row, 5, expected)
        ws.cell(row, 7, "PASS")
        ws.cell(row, 9, "ANDROID_old.mov")
    ws.cell(11, 8, "https://example.com/old-video")
    ws.cell(11, 8).hyperlink = "https://example.com/old-video"
    ws.merge_cells("H11:H13")
    ws.cell(14, 11, "Known flaky on Android")
    ws.column_dimensions["D"].width = 41.25
    ws.column_dimensions["E"].width = 55.88
    ws.freeze_panes = "A11"
    dv = DataValidation(type="list", formula1=dv_values, allow_blank=True)
    dv.add("G11:G1000")
    ws.add_data_validation(dv)
    ws.conditional_formatting.add("G11:G1000", CellIsRule(operator="equal", formula=['"PASS"'],
                                  fill=PatternFill("solid", start_color="FFB7E1CD", end_color="FFB7E1CD")))


def make_workbook(path: Path) -> Path:
    wb = Workbook()
    history = wb.active
    history.title = "1712"
    make_sheet(history, with_result_header=False)
    history.sheet_state = "hidden"
    make_sheet(wb.create_sheet("app"))
    wb.active = 0          # Google exports leave the (hidden) first tab active ...
    wb.views = []          # ... and write no <bookViews> at all
    wb.save(path)
    return path


def make_csv(path: Path) -> Path:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(["Name", "Demo Smoke Test Checklist ", "", "", "", "", "", "", "", "", "", ""])
        writer.writerow(["Test cases Version", "v1.0 Checklist", "", ""])
        writer.writerow(["Sprint", "", "All Cases", "0"])
        writer.writerow(["Build version", "", "", ""])
        writer.writerow(["Pre-condition ", "", "", ""])
        writer.writerow(["Pass Cases ", "6", "Unavailabe Cases", "0"])
        writer.writerow(["", "", "Fail Cases", "0"])
        writer.writerow(["Tested by", "QA", "", ""])
        writer.writerow(["Test Date", "Apr 02, 2026", "", ""])
        writer.writerow(HEADER)
        last_section = None
        for platform, section, desc, expected in CASES:
            last_section = section or last_section
            writer.writerow(["", platform, section or "", desc, expected, "", "PASS", "", "", "", "", ""])
    return path


def fake_run(cases_doc, platform, statuses, run_dir: Path) -> dict:
    results = {}
    for case, status in zip(cases_doc["cases"], statuses):
        if status is None:
            continue
        rec = {"id": case["id"], "section": case["section"], "row": case["row"], "status": status,
               "message": "" if status == "pass" else f"{status} message for {case['id']}", "duration_s": 12.0}
        if status in ("pass", "fail"):
            evidence = run_dir / "evidence" / platform / f"{case['id'].lower()}.png"
            evidence.parent.mkdir(parents=True, exist_ok=True)
            evidence.write_bytes(b"png")
            rec["evidence"] = evidence.relative_to(run_dir).as_posix()
        results[case["key"]] = rec
    return {"platform": platform, "started_at": "2026-10-27T10:00:00+07:00",
            "app": {"version": "1.2.3", "build": "45"}, "results": results}


class ParseTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_xlsx_picks_visible_tab_and_maps_columns(self):
        doc = testcases.load_cases(make_workbook(self.dir / "Demo.xlsx"))
        src = doc["source"]
        self.assertEqual(src["sheet"], "app")
        self.assertEqual(src["header_row"], 10)
        self.assertEqual(src["columns"]["result"], 7)
        self.assertEqual(src["columns"]["ios"], 8)
        self.assertEqual(src["columns"]["note"], 11)
        self.assertEqual(len(doc["cases"]), len(CASES))
        self.assertEqual([c["id"] for c in doc["cases"][:3]], ["TC-001", "TC-002", "TC-003"])
        self.assertEqual(doc["cases"][1]["section"], "Registration")  # carried forward
        self.assertEqual(doc["cases"][3]["platforms"], ["ios"])
        self.assertEqual(doc["cases"][4]["platforms"], ["android"])

    def test_hidden_tab_without_result_header_uses_validation_column(self):
        doc = testcases.load_cases(make_workbook(self.dir / "Demo.xlsx"), sheet="1712")
        self.assertEqual(doc["source"]["columns"]["result"], 7)

    def test_csv_and_xlsx_produce_the_same_keys(self):
        x = testcases.load_cases(make_workbook(self.dir / "Demo.xlsx"))
        c = testcases.load_cases(make_csv(self.dir / "Demo - app.csv"))
        self.assertEqual(c["source"]["sheet"], "app")
        self.assertEqual([k["key"] for k in x["cases"]], [k["key"] for k in c["cases"]])

    def test_header_aliases(self):
        self.assertEqual(testcases.header_field("Evidence Android(dd/mm/yyyy)"), "android")
        self.assertEqual(testcases.header_field("Evidence IOS  (dd/mm/yyyy)"), "ios")
        self.assertEqual(testcases.header_field("Kết quả mong đợi"), "expected")
        self.assertEqual(testcases.header_field("Status"), "result")
        self.assertIsNone(testcases.header_field("Test cases Version"))

    def test_standard_template_roundtrip(self):
        generated = [{"section": "Login", "description": "Login with valid credentials", "expected": "Home shown"},
                     {"section": "Login", "description": "Wrong password", "expected": "Error toast"}]
        path = template.write_template(self.dir / "gen.xlsx", "Demo Smoke Test Checklist", generated)
        doc = testcases.load_cases(path)
        self.assertEqual(doc["source"]["header_row"], template.STANDARD_HEADER_ROW)
        self.assertEqual(len(doc["cases"]), 2)
        ws = load_workbook(path)["Test cases"]
        self.assertEqual(ws["A1"].fill.fgColor.rgb, template.BLUE)
        self.assertIsNone(ws["B1"].fill.fgColor.rgb if ws["B1"].fill.fill_type else None)
        self.assertTrue(ws["B6"].value.startswith("=COUNTIF(G11:G12"))


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.run_dir = self.dir / "runs" / "2026-10-27_100000"
        self.run_dir.mkdir(parents=True)

    def tearDown(self):
        self.tmp.cleanup()

    def test_xlsx_export_adds_dated_tab_with_identical_layout(self):
        source = make_workbook(self.dir / "Demo.xlsx")
        doc = testcases.load_cases(source)
        run = fake_run(doc, "ios", ["pass", "fail", "blocked", "manual", "skipped", None], self.run_dir)
        out = report.export(doc, [run], self.run_dir, tester="Bot", run_date=dt.date(2026, 10, 27))
        self.assertTrue(out["output"].endswith("Demo - 2710.xlsx"))
        self.assertEqual(out["counts"], {"PASS": 1, "FAIL": 1, "PENDING": 4})
        self.assertEqual(report.verify(Path(out["output"]), out["sheet"], doc), out["counts"])

        wb = load_workbook(out["output"])
        self.assertEqual(wb.sheetnames, ["1712", "app", "2710"])
        self.assertEqual(wb.active.title, "2710")
        new, src = wb["2710"], wb["app"]
        self.assertEqual(new.sheet_state, "visible")
        self.assertEqual(new.freeze_panes, "A11")
        self.assertEqual(new.column_dimensions["E"].width, src.column_dimensions["E"].width)
        self.assertEqual(new["A1"].fill.fgColor.rgb, "FF3C78D8")
        # results
        self.assertEqual([new.cell(r, 7).value for r in range(11, 17)], ["PASS", "FAIL", "PENDING", "PENDING", "PENDING", "PENDING"])
        self.assertEqual(new["F11"].value, "As expected")
        self.assertIn("fail message", new["F12"].value)
        self.assertEqual(new["A11"].value, "TC-001")
        self.assertEqual(new["H11"].value, "evidence/ios/tc-001.png")
        self.assertEqual(new["H11"].hyperlink.target, "evidence/ios/tc-001.png")
        self.assertIsNone(new["I11"].value)                      # old Android evidence cleared
        self.assertIn("[Blocked]", new["K13"].value)
        self.assertTrue(new["K14"].value.startswith("Known flaky on Android\n[Manual]"))
        self.assertEqual(new["L11"].value, "Maestro (auto)")
        self.assertIsNone(new["L13"].value)
        # evidence merge split, summary merges kept
        merged = {str(m) for m in new.merged_cells.ranges}
        self.assertNotIn("H11:H13", merged)
        self.assertIn("B8:D8", merged)
        # dropdown extended with PENDING, colouring copied
        self.assertIn("PENDING", new.data_validations.dataValidation[0].formula1)
        self.assertEqual(len(list(new.conditional_formatting)), 1)
        # summary block
        self.assertEqual(new["B8"].value, "Bot")
        self.assertEqual(new["B9"].value, dt.datetime(2026, 10, 27))
        self.assertEqual(new["B4"].value, "iOS 1.2.3 (45)")
        self.assertEqual(new["B6"].value, '=COUNTIF(G11:G16,"PASS")')
        self.assertEqual(new["D6"].value, '=COUNTIF(G11:G16,"PENDING")')
        self.assertEqual(new["D7"].value, '=COUNTIF(G11:G16,"FAIL")')
        self.assertEqual(new["D3"].value, "=COUNTA(D11:D16)")
        # the source tab is untouched
        self.assertEqual(src["G12"].value, "PASS")
        self.assertIn("H11:H13", {str(m) for m in src.merged_cells.ranges})

    def test_second_export_same_day_gets_unique_tab(self):
        source = make_workbook(self.dir / "Demo.xlsx")
        doc = testcases.load_cases(source)
        run = fake_run(doc, "ios", ["pass"] * 6, self.run_dir)
        first = report.export(doc, [run], self.run_dir, run_date=dt.date(2026, 10, 27))
        doc2 = dict(doc, source=dict(doc["source"], path=first["output"]))
        second = report.export(doc2, [run], self.run_dir, out_path=self.dir / "again.xlsx",
                               run_date=dt.date(2026, 10, 27))
        self.assertEqual(second["sheet"], "2710 (2)")

    def test_csv_export_and_only_new_tab(self):
        source = make_csv(self.dir / "Demo - app.csv")
        doc = testcases.load_cases(source)
        run = fake_run(doc, "android", ["pass", "pass", "fail", "pass", "pass", "error"], self.run_dir)
        out = report.export(doc, [run], self.run_dir, run_date=dt.date(2026, 10, 27), only_new_tab=True)
        wb = load_workbook(out["output"])
        self.assertEqual(wb.sheetnames, ["2710"])
        ws = wb["2710"]
        self.assertEqual(ws["G11"].value, "PASS")
        self.assertEqual(ws["I11"].value, "evidence/android/tc-001.png")   # Android evidence column
        self.assertEqual(ws["G16"].value, "PENDING")
        self.assertIn("[Automation error]", ws["K16"].value)
        self.assertEqual(ws["A1"].fill.fgColor.rgb, template.BLUE)
        self.assertEqual(ws["A10"].fill.fgColor.rgb, template.GREY)
        self.assertEqual(ws.freeze_panes, "A11")
        self.assertTrue(any("PENDING" in (dv.formula1 or "") for dv in ws.data_validations.dataValidation))

    def test_combine_multi_platform(self):
        case = {"key": "k1", "id": "TC-001"}
        ios = {"platform": "ios", "results": {"k1": {"status": "pass", "evidence": "evidence/ios/a.png"}}}
        android_fail = {"platform": "android", "results": {"k1": {"status": "fail", "message": "boom"}}}
        android_skip = {"platform": "android", "results": {"k1": {"status": "skipped"}}}
        android_missing = {"platform": "android", "results": {}}
        partial = {"platform": "ios", "results": {"k1": {"status": "pass", "coverage": "partial", "uncovered": "email"}}}
        self.assertEqual(report.combine(case, [ios]).status, "PASS")
        verdict = report.combine(case, [ios, android_fail])
        self.assertEqual(verdict.status, "FAIL")
        self.assertIn("Android: boom", verdict.actual)
        self.assertEqual(report.combine(case, [ios, android_skip]).status, "PASS")
        self.assertEqual(report.combine(case, [ios, android_missing]).status, "PENDING")
        verdict = report.combine(case, [partial])
        self.assertEqual(verdict.status, "PENDING")
        self.assertIn("verify manually: email", verdict.note)
        manual = {"status": "manual", "message": "Needs the OTP email"}
        verdict = report.combine(case, [{"platform": "ios", "results": {"k1": manual}},
                                        {"platform": "android", "results": {"k1": dict(manual)}}])
        self.assertEqual(verdict.note, "[Manual] Needs the OTP email")          # said once, no platform prefix


class UploadTests(unittest.TestCase):
    """Network is faked: upload._post is replaced, nothing leaves the machine."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.run_dir = Path(self.tmp.name)
        for rel, data in (("evidence/ios/tc-001.png", b"PNGIOS"), ("evidence/android/tc-001.png", b"PNGAND"),
                          ("evidence/android/tc-001.mp4", b"MP4")):
            (self.run_dir / rel).parent.mkdir(parents=True, exist_ok=True)
            (self.run_dir / rel).write_bytes(data)
        self.runs = [
            {"platform": "ios", "results": {"k1": {"status": "pass", "evidence": "evidence/ios/tc-001.png"}}},
            {"platform": "android", "results": {"k1": {"status": "pass", "evidence": "evidence/android/tc-001.png",
                                                       "video": "evidence/android/tc-001.mp4"}}},
        ]
        self.calls = []

        def fake_post(url, body, content_type, timeout=180, headers=None):
            self.calls.append((url, body, content_type))
            if "imgbb" in url:
                return json.dumps({"success": True, "data": {"url": f"https://i.ibb.co/{len(self.calls)}.png"}}).encode()
            return f"https://files.catbox.moe/f{len(self.calls)}.png".encode()

        patches = [mock.patch.object(upload, "_post", fake_post), mock.patch.object(upload.time, "sleep")]
        for p in patches:
            p.start()
            self.addCleanup(p.stop)

    def tearDown(self):
        self.tmp.cleanup()

    def test_multipart_body(self):
        body, ctype = upload._multipart({"reqtype": "fileupload"}, "fileToUpload", self.run_dir / "evidence/ios/tc-001.png")
        boundary = ctype.split("boundary=")[1]
        self.assertTrue(body.startswith(f"--{boundary}\r\n".encode()))
        self.assertIn(b'name="reqtype"\r\n\r\nfileupload\r\n', body)
        self.assertIn(b'filename="tc-001.png"\r\nContent-Type: image/png\r\n\r\nPNGIOS', body)
        self.assertTrue(body.endswith(f"\r\n--{boundary}--\r\n".encode()))

    def test_upload_is_cached_per_file_content(self):
        links = upload.upload_evidence(self.run_dir, self.runs, "catbox")
        self.assertEqual(len(links), 3)
        self.assertTrue(all(url == "https://catbox.moe/user/api.php" for url, _, _ in self.calls))
        self.assertEqual(upload.upload_evidence(self.run_dir, self.runs, "catbox"), links)
        self.assertEqual(len(self.calls), 3)                                  # nothing uploaded twice
        (self.run_dir / "evidence/ios/tc-001.png").write_bytes(b"NEW")
        upload.upload_evidence(self.run_dir, self.runs, "catbox")
        self.assertEqual(len(self.calls), 4)                                  # only the changed file

    def test_plan_has_no_side_effects(self):
        plan = {p["file"]: p["action"] for p in upload.plan_uploads(self.run_dir, self.runs, "imgbb")}
        self.assertEqual(plan["evidence/ios/tc-001.png"], "upload")
        self.assertEqual(plan["evidence/android/tc-001.mp4"], "skip (images only)")
        self.assertEqual(self.calls, [])
        upload.upload_evidence(self.run_dir, self.runs, "catbox")
        plan = {p["file"]: p["action"] for p in upload.plan_uploads(self.run_dir, self.runs, "catbox")}
        self.assertTrue(plan["evidence/ios/tc-001.png"].startswith("reuse https://files.catbox.moe/"))

    def test_host_rules(self):
        upload.upload_evidence(self.run_dir, self.runs, "litterbox", expire="24h")
        self.assertIn(b'name="time"\r\n\r\n24h\r\n', self.calls[0][1])
        cache = json.loads((self.run_dir / "uploads.json").read_text())
        self.assertTrue(all(entry["expires_at"] for entry in cache["files"].values()))
        with self.assertRaises(UserError):
            upload.upload_evidence(self.run_dir, self.runs, "litterbox", expire="7d")
        with self.assertRaises(UserError):
            upload.upload_evidence(self.run_dir, self.runs, "catbox", expire="1h")
        with mock.patch.dict(os.environ, {}, clear=False):
            os.environ.pop("IMGBB_API_KEY", None)
            with self.assertRaises(UserError):
                upload.upload_evidence(self.run_dir, self.runs, "imgbb")
        with mock.patch.dict(os.environ, {"IMGBB_API_KEY": "test-key"}):
            links = upload.upload_evidence(self.run_dir, self.runs, "imgbb")
        self.assertEqual(sorted(links), ["evidence/android/tc-001.png", "evidence/ios/tc-001.png"])  # no video

    def test_droplr(self):
        seen = []

        def droplr_post(url, body, content_type, timeout=180, headers=None):
            seen.append((url, body, content_type, headers))
            return json.dumps({"code": f"c{len(seen)}", "shortlink": f"http://d.pr/i/c{len(seen)}"}).encode()

        creds = {"DROPLR_USERNAME": "qa@example.com", "DROPLR_PASSWORD": "s3cret", "DROPLR_BOARD": "b42"}
        with mock.patch.dict(os.environ, creds), mock.patch.object(upload, "_post", side_effect=droplr_post):
            links = upload.upload_evidence(self.run_dir, self.runs, "droplr")
        self.assertEqual(len(links), 3)                                      # images + the video
        self.assertTrue(all(url.startswith("https://d.pr/i/") for url in links.values()))
        url, body, ctype, headers = next(s for s in seen if "tc-001.png" in s[0] and s[1] == b"PNGIOS")
        self.assertEqual(url, "https://api.droplr.com/files?filename=tc-001.png&board=b42")
        self.assertEqual(ctype, "image/png")
        self.assertEqual(headers["Authorization"], "Basic cWFAZXhhbXBsZS5jb206czNjcmV0")   # qa@example.com:s3cret
        self.assertTrue(any(s[2] == "video/mp4" for s in seen))

        with mock.patch.dict(os.environ, {"DROPLR_TOKEN": "jwt-123"}), \
                mock.patch.object(upload, "_post", side_effect=droplr_post):
            upload.upload_evidence(self.run_dir, self.runs, "droplr", force=True)
        self.assertEqual(seen[-1][3]["Authorization"], "Bearer jwt-123")

    def test_droplr_errors(self):
        env = {k: v for k, v in os.environ.items() if not k.startswith("DROPLR_")}
        with mock.patch.dict(os.environ, env, clear=True):
            with self.assertRaises(UserError):                               # no credentials
                upload.upload_evidence(self.run_dir, self.runs, "droplr")
        denied = urllib.error.HTTPError("https://api.droplr.com/files", 401, "Unauthorized", {},
                                        io.BytesIO(b'{"error": "bad credentials"}'))
        self.addCleanup(denied.close)
        with mock.patch.dict(os.environ, {"DROPLR_TOKEN": "bad"}), \
                mock.patch.object(upload, "_post", side_effect=denied):
            with self.assertRaises(UserError):                               # wrong credentials: stop at once
                upload.upload_evidence(self.run_dir, self.runs, "droplr")
        with mock.patch.dict(os.environ, {"DROPLR_TOKEN": "ok"}), \
                mock.patch.object(upload, "_post", return_value=b'{"code": "x1"}'):
            self.assertEqual(upload.upload_evidence(self.run_dir, self.runs, "droplr", retries=0), {})

    def test_failures_keep_local_links(self):
        with mock.patch.object(upload, "_post", side_effect=urllib.error.URLError("offline")):
            self.assertEqual(upload.upload_evidence(self.run_dir, self.runs, "catbox", retries=1), {})
        with mock.patch.object(upload, "_post", return_value=b"Internal error, try again"):
            self.assertEqual(upload.upload_evidence(self.run_dir, self.runs, "catbox", retries=0), {})
        self.assertFalse((self.run_dir / "uploads.json").exists())

    def test_export_writes_public_links(self):
        source = make_workbook(self.run_dir / "Demo.xlsx")
        doc = testcases.load_cases(source)
        run = fake_run(doc, "ios", ["pass", "fail", "manual", None, None, None], self.run_dir)
        links = {run["results"][doc["cases"][0]["key"]]["evidence"]: "https://files.catbox.moe/abc.png"}
        out = report.export(doc, [run], self.run_dir, run_date=dt.date(2026, 10, 27), links=links)
        ws = load_workbook(out["output"])["2710"]
        self.assertEqual(ws["H11"].value, "https://files.catbox.moe/abc.png")
        self.assertEqual(ws["H11"].hyperlink.target, "https://files.catbox.moe/abc.png")
        self.assertEqual(ws["H12"].value, "evidence/ios/tc-002.png")          # not uploaded -> local link
        self.assertEqual(out["evidence_links"], {"uploaded": 1, "local": 1})


class FakeGoogle:
    """Minimal stand-in for the OAuth token endpoint and the Drive v3 API."""

    def __init__(self, email="qa@company.com"):
        self.email, self.calls, self.folders, self.n = email, [], {}, 0

    def __call__(self, method, url, body=None, headers=None, timeout=180):
        import re
        import urllib.parse as up

        self.calls.append((method, url, body, headers or {}))
        parsed, query = up.urlparse(url), dict(up.parse_qsl(up.urlparse(url).query))
        if url.startswith(gdrive.TOKEN_URL):
            token = {"access_token": "at-1", "expires_in": 3600}
            if b"grant_type=authorization_code" in (body or b""):
                token["refresh_token"] = "rt-new"
            return 200, {}, json.dumps(token).encode()
        if parsed.path.endswith("/about"):
            return 200, {}, json.dumps({"user": {"emailAddress": self.email}}).encode()
        if parsed.path == "/drive/v3/files" and method == "GET":
            name = re.search(r"name = '((?:[^'\\]|\\.)*)'", query["q"]).group(1).replace("\\'", "'")
            parent = re.search(r"'([^']+)' in parents", query["q"]).group(1)
            fid = self.folders.get((name, parent))
            files = [{"id": fid, "webViewLink": f"https://drive.google.com/drive/folders/{fid}"}] if fid else []
            return 200, {}, json.dumps({"files": files}).encode()
        if parsed.path == "/drive/v3/files" and method == "POST":
            meta = json.loads(body)
            self.n += 1
            self.folders[(meta["name"], (meta.get("parents") or ["root"])[0])] = f"fold{self.n}"
            return 200, {}, json.dumps({"id": f"fold{self.n}",
                                        "webViewLink": f"https://drive.google.com/drive/folders/fold{self.n}"}).encode()
        if "/permissions" in parsed.path:
            return 200, {}, b'{"id": "perm1"}'
        if parsed.path == "/upload/drive/v3/files" and query.get("uploadType") == "resumable":
            self.n += 1
            return 200, {"location": f"https://www.googleapis.com/upload/drive/v3/files?upload_id=u{self.n}"}, b""
        if method == "PUT" and "upload_id=" in url:
            fid = "file" + url.split("upload_id=")[1]
            return 200, {}, json.dumps({"id": fid, "webViewLink": f"https://drive.google.com/file/d/{fid}/view"}).encode()
        return 404, {}, b'{"error": "not found"}'


class GDriveTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        home = Path(self.tmp.name) / "home"
        self.run_dir = Path(self.tmp.name) / "runs" / "2026-10-27_140158"
        for rel in ("evidence/ios/tc-004.png", "evidence/android/tc-004.png", "evidence/android/tc-004.mp4"):
            (self.run_dir / rel).parent.mkdir(parents=True, exist_ok=True)
            (self.run_dir / rel).write_bytes(rel.encode())
        (self.run_dir / "cases.json").write_text(json.dumps({"source": {"path": "/x/[App - D1] Smoke testing.xlsx"}}))
        self.runs = [{"platform": "ios", "results": {"k": {"evidence": "evidence/ios/tc-004.png"}}},
                     {"platform": "android", "results": {"k": {"evidence": "evidence/android/tc-004.png",
                                                               "video": "evidence/android/tc-004.mp4"}}}]
        self.google = FakeGoogle()
        for patcher in (mock.patch.object(gdrive, "HOME", home),
                        mock.patch.object(gdrive, "TOKEN_FILE", home / "google-token.json"),
                        mock.patch.object(gdrive, "CLIENT_FILE", home / "google-oauth-client.json"),
                        mock.patch.object(gdrive, "BUNDLED_CLIENT", Path(self.tmp.name) / "bundled-client.json"),
                        mock.patch.object(gdrive, "RESOURCE_DIR", Path(self.tmp.name) / "resource"),
                        mock.patch.object(gdrive, "_http", self.google),
                        mock.patch.dict(os.environ, {"RNSMOKE_GOOGLE_CLIENT_ID": "cid",
                                                     "RNSMOKE_GOOGLE_CLIENT_SECRET": "csecret"})):
            patcher.start()
            self.addCleanup(patcher.stop)
        gdrive._save_token({"refresh_token": "rt-1", "access_token": None, "expires_at": 0})

    def tearDown(self):
        self.tmp.cleanup()

    def test_pkce_and_auth_url(self):
        import base64 as b64
        import hashlib as hl

        verifier, challenge = gdrive.pkce_pair()
        self.assertEqual(challenge, b64.urlsafe_b64encode(hl.sha256(verifier.encode()).digest()).decode().rstrip("="))
        url = gdrive.auth_url("cid", "http://127.0.0.1:5000/", challenge, "st")
        for part in ("scope=https%3A%2F%2Fwww.googleapis.com%2Fauth%2Fdrive.file", "code_challenge_method=S256",
                     "access_type=offline", "redirect_uri=http%3A%2F%2F127.0.0.1%3A5000%2F"):
            self.assertIn(part, url)

    def test_client_lookup_and_import(self):
        no_env = {"RNSMOKE_GOOGLE_CLIENT_ID": "", "RNSMOKE_GOOGLE_CLIENT_SECRET": "", "RNSMOKE_GOOGLE_CLIENT_FILE": ""}
        with mock.patch.dict(os.environ, no_env):
            with self.assertRaises(UserError):
                gdrive.load_client()
            gdrive.BUNDLED_CLIENT.write_text(json.dumps({"installed": {"client_id": "bundled-id",
                                                                       "client_secret": "bundled-secret"}}))
            self.assertEqual(gdrive.load_client(), ("bundled-id", "bundled-secret"))      # shipped with the repo
            downloaded = Path(self.tmp.name) / "client_secret_123.json"
            downloaded.write_text(json.dumps({"installed": {"client_id": "own-id", "client_secret": "own-secret"}}))
            stored = gdrive.import_client(downloaded)
            self.assertEqual(oct(stored.stat().st_mode & 0o777), "0o600")
            self.assertEqual(gdrive.load_client(), ("own-id", "own-secret"))              # personal file wins
            downloaded.write_text('{"not": "a client"}')
            with self.assertRaises(UserError):
                gdrive.import_client(downloaded)

    def test_link_from_id_and_downloaded_client(self):
        def no_view_link(method, url, body=None, headers=None, timeout=180):
            if method == "PUT":
                return 200, {}, b'{"id": "abc123"}'
            return self.google(method, url, body, headers, timeout)

        with mock.patch.object(gdrive, "_http", side_effect=no_view_link):
            link = gdrive.upload_file(self.run_dir / "evidence/ios/tc-004.png", "fold", "ios-tc-004.png", "image/png")
        self.assertEqual(link, "https://drive.google.com/file/d/abc123/view")
        downloads = Path(self.tmp.name) / "Downloads"
        downloads.mkdir()
        (downloads / "client_secret_1-abc.apps.googleusercontent.com.json").write_text("{}")
        with mock.patch.object(gdrive.Path, "home", return_value=Path(self.tmp.name)):
            self.assertEqual(gdrive.find_downloaded_client().name, "client_secret_1-abc.apps.googleusercontent.com.json")

    def test_browser_login_via_loopback(self):
        import threading
        import urllib.parse as up
        import urllib.request as ur

        def fake_browser(url):
            query = dict(up.parse_qsl(up.urlparse(url).query))
            redirect = f"{query['redirect_uri']}?code=auth-code&state={query['state']}"
            threading.Thread(target=lambda: ur.urlopen(redirect, timeout=10).read(), daemon=True).start()
            return True

        gdrive.TOKEN_FILE.unlink()
        with mock.patch.object(gdrive.webbrowser, "open", side_effect=fake_browser):
            record = gdrive.login(timeout=20)
        self.assertEqual(record["refresh_token"], "rt-new")
        self.assertEqual(record["account"], "qa@company.com")
        exchange = next(c for c in self.google.calls if c[2] and b"grant_type=authorization_code" in c[2])
        self.assertIn(b"code=auth-code", exchange[2])
        self.assertIn(b"code_verifier=", exchange[2])
        self.assertEqual(oct(gdrive.TOKEN_FILE.stat().st_mode & 0o777), "0o600")

    def test_private_upload_and_cache(self):
        links = upload.upload_evidence(self.run_dir, self.runs, "gdrive")
        self.assertEqual(len(links), 3)
        self.assertTrue(all(u.startswith("https://drive.google.com/file/d/") for u in links.values()))
        self.assertEqual(sorted(self.google.folders), [("2026-10-27_140158", "fold2"),
                                                       ("RN smoke evidence", "root"),
                                                       ("[App - D1] Smoke testing", "fold1")])
        self.assertFalse(any("/permissions" in c[1] for c in self.google.calls))        # private by default
        upload_meta = [json.loads(c[2]) for c in self.google.calls if "uploadType=resumable" in c[1]]
        self.assertIn({"name": "ios-tc-004.png", "parents": ["fold3"]}, upload_meta)
        self.assertTrue(all(c[3].get("Authorization") == "Bearer at-1" for c in self.google.calls
                            if "googleapis.com/drive" in c[1] or "uploadType" in c[1]))
        token = json.loads(gdrive.TOKEN_FILE.read_text())
        self.assertEqual(token["access_token"], "at-1")
        self.assertEqual(oct(gdrive.TOKEN_FILE.stat().st_mode & 0o777), "0o600")
        before = len(self.google.calls)
        self.assertEqual(upload.upload_evidence(self.run_dir, self.runs, "gdrive"), links)
        self.assertEqual(len(self.google.calls), before)                                # cached: no network

    def test_share_modes(self):
        upload.upload_evidence(self.run_dir, self.runs, "gdrive", share="domain")
        permission = next(json.loads(c[2]) for c in self.google.calls if "/permissions" in c[1])
        self.assertEqual(permission, {"type": "domain", "role": "reader", "domain": "company.com"})
        self.google.email = "someone@gmail.com"
        with self.assertRaises(UserError):
            upload.upload_evidence(self.run_dir, self.runs, "gdrive", share="domain", force=True)
        with self.assertRaises(UserError):
            upload.upload_evidence(self.run_dir, self.runs, "catbox", share="domain")

    def test_setup_problems(self):
        with mock.patch.dict(os.environ, {"RNSMOKE_GOOGLE_CLIENT_ID": "", "RNSMOKE_GOOGLE_CLIENT_SECRET": ""}):
            with self.assertRaises(UserError):
                gdrive.load_client()
        gdrive.TOKEN_FILE.unlink()
        with self.assertRaises(UserError):
            gdrive.access_token(interactive=False)
        gdrive._save_token({"refresh_token": "rt-1", "expires_at": 0})
        with mock.patch.object(gdrive, "_http", side_effect=urllib.error.URLError("offline")):
            self.assertEqual(upload.upload_evidence(self.run_dir, self.runs, "gdrive"), {})   # local links kept


class FakeRclone:
    """Stands in for the rclone binary (upload._rclone_cmd)."""

    def __init__(self):
        self.calls, self.files, self.ready = [], {}, True

    def __call__(self, args, timeout=300):
        self.calls.append(list(args))
        cmd = args[0]
        if cmd == "config" and args[1] == "dump":
            remote = {"type": "drive", "scope": "drive.file"}
            if self.ready:
                remote["token"] = '{"access_token":"x"}'
            return 0, json.dumps({"rnsmoke-gdrive": remote}), ""
        if cmd == "mkdir":
            return 0, "", ""
        if cmd == "copyto":
            self.files.setdefault(args[2], f"fileid{len(self.files) + 1}")    # overwrite keeps the same id
            return 0, "", ""
        if cmd == "lsjson":
            target = args[-1]
            if target in self.files:
                return 0, json.dumps({"ID": self.files[target], "IsDir": False}), ""
            return 0, json.dumps({"ID": "folderid", "IsDir": True}), ""
        if cmd == "link":
            return 0, "https://drive.google.com/drive/folders/folderid?usp=sharing\n", ""
        return 1, "", f"unexpected {cmd}"


class RcloneTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.run_dir = Path(self.tmp.name) / "runs" / "2026-10-27_140158"
        for rel in ("evidence/ios/tc-004.png", "evidence/android/tc-004.png"):
            (self.run_dir / rel).parent.mkdir(parents=True, exist_ok=True)
            (self.run_dir / rel).write_bytes(rel.encode())
        (self.run_dir / "cases.json").write_text(json.dumps({"source": {"path": "/x/[App - D1] Smoke testing.xlsx"}}))
        self.runs = [{"platform": "ios", "results": {"k": {"evidence": "evidence/ios/tc-004.png"}}},
                     {"platform": "android", "results": {"k": {"evidence": "evidence/android/tc-004.png"}}}]
        self.rclone = FakeRclone()
        self.real_rclone_cmd = upload._rclone_cmd
        patcher = mock.patch.object(upload, "_rclone_cmd", self.rclone)
        patcher.start()
        self.addCleanup(patcher.stop)

    def tearDown(self):
        self.tmp.cleanup()

    def test_private_upload_with_drive_links(self):
        links = upload.upload_evidence(self.run_dir, self.runs, "rclone")
        self.assertEqual(sorted(links.values()), ["https://drive.google.com/file/d/fileid1/view",
                                                  "https://drive.google.com/file/d/fileid2/view"])
        folder = "rnsmoke-gdrive:RN smoke evidence/[App - D1] Smoke testing/2026-10-27_140158"
        self.assertIn(["mkdir", folder], self.rclone.calls)
        self.assertIn(["copyto", str(self.run_dir / "evidence/ios/tc-004.png"), f"{folder}/ios-tc-004.png"],
                      self.rclone.calls)
        self.assertFalse(any(c[0] == "link" for c in self.rclone.calls))            # private by default
        calls = len(self.rclone.calls)
        upload.upload_evidence(self.run_dir, self.runs, "rclone")
        self.assertEqual(len(self.rclone.calls), calls)                             # cached
        (self.run_dir / "evidence/ios/tc-004.png").write_bytes(b"retaken")
        again = upload.upload_evidence(self.run_dir, self.runs, "rclone")
        self.assertEqual(again, links)                                              # same path -> same link

    def test_share_and_setup_errors(self):
        upload.upload_evidence(self.run_dir, self.runs, "rclone", share="anyone")
        self.assertTrue(any(c[0] == "link" for c in self.rclone.calls))
        with self.assertRaises(UserError):
            upload.upload_evidence(self.run_dir, self.runs, "rclone", share="domain", force=True)
        self.rclone.ready = False
        with self.assertRaises(UserError):                                          # rclone-setup not done
            upload.upload_evidence(self.run_dir, self.runs, "rclone", force=True)

    def test_missing_binary(self):
        with mock.patch.object(upload.shutil, "which", return_value=None):
            with self.assertRaises(UserError):
                self.real_rclone_cmd(["listremotes"])


class SafetyTests(unittest.TestCase):
    """Data-safety guard: one session account, no new accounts or records."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        (self.root / "subflows").mkdir()
        (self.root / "flows").mkdir()

    def tearDown(self):
        self.tmp.cleanup()

    def flow(self, name, body):
        path = self.root / name
        path.write_text("appId: ${APP_ID}\n---\n" + body, encoding="utf-8")
        return path

    def rules(self, path, level="none"):
        return sorted({f.rule for f in safety.violations(safety.check_flow(path, level))})

    def test_safe_login_flow_passes(self):
        path = self.flow("flows/login.yaml", """
- launchApp:
    clearState: true
- tapOn: "Email address"
- inputText: "${TEST_EMAIL}"
- tapOn:
    text: "Password"
    below:
      text: "Sign up today"          # anchor text is not a tap target
- inputText: ${TEST_PASSWORD}
- tapOn: "Sign in"
- extendedWaitUntil:
    visible: "Create account"         # assertions are never taps
    timeout: 10000
- takeScreenshot: evidence
""")
        self.assertEqual(self.rules(path), [])

    def test_new_identities_are_refused(self):
        for body in ('- inputText: "qa+${new Date().getTime()}@example.com"\n',
                     '- inputText: "someone@company.com"\n',
                     "- inputRandomEmail\n",
                     "- inputRandomPersonName\n",
                     '- inputText: "${TEST_EMAIL_2}"\n',
                     '- inputText: ${ADMIN_PASSWORD}\n',
                     "- runFlow:\n    file: ../subflows/x.yaml\n    env:\n      PASSWORD: \"Secret#1\"\n"):
            with self.subTest(body=body):
                (self.root / "subflows/x.yaml").write_text("appId: x\n---\n- inputText: ${PASSWORD}\n")
                self.assertIn("identity", self.rules(self.flow("flows/f.yaml", body)))
        self.assertEqual(self.rules(self.flow("flows/ok.yaml", '- inputText: "P@ssw0rd!"\n')), [])  # not an e-mail

    def test_account_actions_are_hard_blocked(self):
        for target in ("CREATE ACCOUNT", "(?i)sign up", "Register", "Delete my account", "Add new profile",
                       "Change password", "Reset password"):
            with self.subTest(target=target):
                path = self.flow("flows/f.yaml", f'- tapOn: "{target}"\n# rnsmoke-safe-tap: "{target}" -- try to excuse it\n')
                self.assertEqual(self.rules(path, "account"), ["account"])          # annotations cannot clear it

    def test_records_edits_and_annotations(self):
        path = self.flow("flows/f.yaml", '- tapOn: "Book"\n- tapOn: "Save"\n- tapOn: "Cancel"\n'
                                         '- tapOn:\n    point: "50%,12%"\n')
        self.assertEqual(self.rules(path), ["edit", "records", "unverifiable"])     # bare "Cancel" is fine
        self.assertEqual(self.rules(path, "account"), ["records", "unverifiable"])  # account level may save
        path = self.flow("flows/f.yaml", '# rnsmoke-safe-tap: "Book" -- bottom tab, navigation only\n'
                                         '# rnsmoke-safe-tap: "point 50%,12%" -- closes the date picker overlay\n'
                                         '- tapOn: "Book"\n- tapOn:\n    point: "50%,12%"\n')
        found = safety.check_flow(path)
        self.assertEqual(safety.violations(found), [])
        self.assertEqual(sorted(f.allowed_by for f in found), ["bottom tab, navigation only",
                                                               "closes the date picker overlay"])

    def test_subflows_env_and_scripts(self):
        (self.root / "subflows/tap.yaml").write_text("appId: x\n---\n- tapOn: ${TARGET}\n")
        (self.root / "subflows/api.js").write_text("http.post('https://api.example.com/users', {body: '{}'})\n")
        path = self.flow("flows/f.yaml", "- runFlow:\n    file: ../subflows/tap.yaml\n    env:\n"
                                         "      TARGET: \"Join waitlist\"\n- runScript: ../subflows/api.js\n"
                                         "- runFlow: ../subflows/missing.yaml\n")
        self.assertEqual(self.rules(path), ["records", "script", "unverifiable"])

    def test_session_account_lock_and_env_filter(self):
        env = safety.filter_env({"TEST_EMAIL": "qa@company.com", "TEST_PASSWORD": "x", "TEST_EMAIL_2": "b@c.com",
                                 "ADMIN_PASSWORD": "y", "TEST_FACILITY": "Demo", "TAP_SPINNER": "z"})
        self.assertEqual(sorted(env), ["TAP_SPINNER", "TEST_EMAIL", "TEST_FACILITY", "TEST_PASSWORD"])
        self.assertEqual(safety.lock_session(self.root, env), "qa***@company.com")
        self.assertEqual(safety.lock_session(self.root, env), "qa***@company.com")    # same account again: fine
        with self.assertRaises(UserError):
            safety.lock_session(self.root, {"TEST_EMAIL": "other@company.com"})
        self.assertNotIn("qa@company.com", (self.root / "session.json").read_text())   # only a fingerprint

    def test_workspace_refuses_unsafe_flows(self):
        (self.root / "package.json").write_text('{"name": "demo"}')
        ws = Workspace(str(self.root))
        ws.ensure()
        ws.save_cases(testcases.load_cases(make_workbook(self.root / "Demo.xlsx")))
        case = select_cases(ws.load_cases()["cases"], ["TC-001"])
        bad = ws.flows_dir / "signup.yaml"
        bad.write_text('appId: x\n---\n- inputText: "a+${new Date().getTime()}@example.com"\n'
                       '- tapOn: "CREATE ACCOUNT"\n- takeScreenshot: evidence\n')
        with self.assertRaises(UserError):
            ws.mark(case, "automated", flow="flows/signup.yaml")
        with self.assertRaises(UserError):
            ws.mark(case, "automated", flow="flows/signup.yaml", requires=["TEST_EMAIL_TWO_PROFILES"])
        good = ws.flows_dir / "home.yaml"
        good.write_text("appId: x\n---\n- launchApp\n- takeScreenshot: evidence\n")
        ws.mark(case, "automated", flow="flows/home.yaml", data="account")
        self.assertEqual(ws.load_manifest()["cases"][case[0]["key"]]["data"], "account")
        bad.write_text(good.read_text() + '- tapOn: "Book now"\n')                  # flow edited later
        ws.mark(case, "automated", flow="flows/home.yaml")
        good.write_text(bad.read_text())
        problems = next(r for r in ws.status_report()["cases"] if r["id"] == "TC-001")["problems"]
        self.assertTrue(any(p.startswith("DATA-SAFETY") for p in problems))
        ws.mark(case, "manual", reason="would create a booking")
        entry = ws.load_manifest()["cases"][case[0]["key"]]
        self.assertNotIn("flow", entry)

    def test_runner_never_hands_an_unsafe_flow_to_maestro(self):
        from rnsmoke.devices import Device

        (self.root / "package.json").write_text('{"name": "demo"}')
        ws = Workspace(str(self.root))
        ws.ensure()
        ws.save_config({"app": {"ios": {"appId": "com.demo"}}, "metro": {"port": 8081}})
        ws.save_cases(testcases.load_cases(make_workbook(self.root / "Demo.xlsx")))
        case = select_cases(ws.load_cases()["cases"], ["TC-001"])
        flow = ws.flows_dir / "f.yaml"
        flow.write_text("appId: x\n---\n- launchApp\n- takeScreenshot: evidence\n")
        ws.mark(case, "automated", flow="flows/f.yaml")
        flow.write_text("appId: x\n---\n- inputRandomEmail\n- tapOn: \"Sign up\"\n")     # edited after marking
        with mock.patch.object(runner, "require_maestro", return_value="maestro"), \
                mock.patch.object(runner.devices, "resolve_device", return_value=Device("ios", "u1", "iPhone", "iOS")), \
                mock.patch.object(runner.devices, "app_info", return_value={"installed": True, "debug": False}), \
                mock.patch.object(runner, "Executor") as executor:
            doc = runner.run(ws, "ios", only=["TC-001"], retries=0)
        executor.return_value.execute.assert_not_called()
        record = doc["results"][case[0]["key"]]
        self.assertEqual(record["status"], "blocked")
        self.assertTrue(record["message"].startswith("Data-safety guard"))


class RunnerHelperTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_junit_and_artifacts(self):
        junit = self.dir / "report.xml"
        junit.write_text("""<?xml version='1.0' encoding='UTF-8'?>
<testsuites><testsuite name="Test Suite" tests="2" failures="1">
  <testcase id="A" name="Login - ok" classname="Login - ok" file="ok.yaml" time="11.4" status="SUCCESS"/>
  <testcase id="B" name="Login - bad" classname="Login - bad" file="bad.yaml" time="6.4" status="ERROR">
    <failure>Assertion is false: "Welcome" is visible</failure>
  </testcase>
</testsuite></testsuites>""", encoding="utf-8")
        reports = runner.parse_junit(junit)
        self.assertEqual([(r["name"], r["failed"]) for r in reports], [("Login - ok", False), ("Login - bad", True)])
        self.assertIn("Welcome", reports[1]["message"])

        flow_dir = self.dir / "2026-10-27_110000" / "Login - bad"
        (flow_dir / "screenshots").mkdir(parents=True)
        (flow_dir / "takeScreenshot").mkdir()
        (flow_dir / "screenshots" / "step-004-assert.png").write_bytes(b"x")
        (flow_dir / "takeScreenshot" / "evidence.png").write_bytes(b"x")
        (flow_dir / "commands.json").write_text(json.dumps([
            {"command": {"tapOnElement": {"selector": {"textRegex": "Sign in"}}}, "metadata": {"status": "COMPLETED", "sequenceNumber": 2}},
            {"command": {"assertConditionCommand": {"condition": {"visible": {"textRegex": "Welcome"}}}},
             "metadata": {"status": "FAILED", "sequenceNumber": 3}},
        ]), encoding="utf-8")
        found = runner.find_artifacts(self.dir, "Login - bad")
        self.assertEqual(found, flow_dir)
        self.assertEqual(runner.pick_screenshot(found, "fail").name, "step-004-assert.png")
        self.assertEqual(runner.pick_screenshot(found, "pass").name, "evidence.png")
        self.assertEqual(runner.describe_failed_step(found), '#4 assertCondition "Welcome"')

    def test_classify_failure(self):
        art = self.dir / "flow"
        (art / "screen-hierarchy").mkdir(parents=True)
        (art / "screen-hierarchy" / "step.json").write_text('{"text": "Could not connect to development server"}')
        self.assertEqual(runner.classify_failure("Element not found", art, "").status, "error")
        self.assertEqual(runner.classify_failure("Unknown Property: timeout at x.yaml", None, "").status, "error")
        self.assertFalse(runner.classify_failure("Unknown Property: timeout", None, "").retryable)
        self.assertEqual(runner.classify_failure('Assertion is false: "x" is visible', None, "").status, "fail")
        self.assertEqual(runner.classify_failure("", None, "Running flow a.yaml\n[Failed] a (3s) (x)").status, "fail")

    def test_flow_name(self):
        flow = self.dir / "f.yaml"
        flow.write_text('appId: ${APP_ID}\nname: "Login - ok"\n---\n- launchApp\n', encoding="utf-8")
        self.assertEqual(runner.flow_name(flow), "Login - ok")
        flow.write_text("appId: x\n---\n- launchApp\n", encoding="utf-8")
        self.assertEqual(runner.flow_name(flow), "f")


class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        (self.dir / "package.json").write_text('{"name": "demo"}', encoding="utf-8")
        self.ws = Workspace(str(self.dir))
        self.ws.ensure()
        self.doc = testcases.load_cases(make_workbook(self.dir / "Demo.xlsx"))
        self.ws.save_cases(self.doc)

    def tearDown(self):
        self.tmp.cleanup()

    def test_select_cases(self):
        cases = self.doc["cases"]
        self.assertEqual([c["id"] for c in select_cases(cases, ["TC-002..TC-004"])], ["TC-002", "TC-003", "TC-004"])
        self.assertEqual(len(select_cases(cases, None, ["login"])), 2)
        self.assertEqual(select_cases(cases, ["row:12"])[0]["id"], "TC-002")
        extra = cases + [dict(cases[-1], id="TC-099", key="k99", section="Home - My Booking")]
        self.assertEqual(len(select_cases(extra, None, ["Home"])), 2)              # exact name wins
        self.assertEqual(len(select_cases(extra, None, ["my booking"])), 1)        # substring fallback

    def test_mark_status_and_relink(self):
        flow = self.ws.flows_dir / "login" / "ok.yaml"
        flow.parent.mkdir(parents=True)
        flow.write_text("appId: ${APP_ID}\n---\n- launchApp\n- takeScreenshot: evidence\n", encoding="utf-8")
        self.ws.mark(select_cases(self.doc["cases"], ["TC-003"]), "automated", flow="flows/login/ok.yaml")
        self.ws.mark(select_cases(self.doc["cases"], ["TC-001"]), "manual", reason="email")
        report_rows = {r["id"]: r for r in self.ws.status_report()["cases"]}
        self.assertEqual(report_rows["TC-003"]["status"], "automated")
        self.assertEqual(report_rows["TC-001"]["reason"], "email")
        self.assertEqual(report_rows["TC-002"]["status"], "todo")

        edited = json.loads(json.dumps(self.doc))
        target = edited["cases"][2]
        target["description"] = "Login with valid credentials."            # small edit -> new key
        target["key"] = testcases.case_key(target["section"], target["description"], target["expected"])
        moved = self.ws.relink(edited["cases"])
        self.assertEqual(moved[0][0], "TC-003")
        self.assertEqual(self.ws.load_manifest()["cases"][target["key"]]["flow"], "flows/login/ok.yaml")

    def test_env_file(self):
        self.ws.env_path.write_text('# comment\nTEST_EMAIL="a@b.c"\nexport TEST_PASSWORD=p=1\nBAD KEY=x\n', encoding="utf-8")
        self.assertEqual(self.ws.env_vars(), {"TEST_EMAIL": "a@b.c", "TEST_PASSWORD": "p=1"})

    def test_init_copies_templates(self):
        self.assertTrue((self.ws.subflows_dir / "launch-clean.yaml").exists())
        self.assertTrue((self.ws.subflows_dir / "hide-keyboard.yaml").exists())
        self.assertTrue((self.ws.subflows_dir / "scroll-until-visible.yaml").exists())
        self.assertIn(".env", (self.ws.root / ".gitignore").read_text())
        self.assertIn("cases.json", (self.ws.root / ".gitignore").read_text())
        self.assertTrue(self.ws.status_report()["warnings"])   # FIRST_SCREEN_TEXT placeholder flagged


if __name__ == "__main__":
    unittest.main()
